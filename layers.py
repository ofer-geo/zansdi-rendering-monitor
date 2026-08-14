from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


EXCEL_PATH = Path("rendering_monitoring_tests.xlsx")
SPECIFICATION_SHEET = "Specification"


def normalize_text(value: object) -> str:
    """Convert an Excel cell value into clean text."""
    if value is None:
        return ""

    return str(value).strip()


def normalize_identifier(value: object) -> str:
    """
    Normalize the GeoServer identifier copied from Excel.

    This also removes an accidental backslash before a colon,
    for example: geonode\\:regions -> geonode:regions
    """
    return normalize_text(value).replace("\\:", ":")


def read_specification() -> tuple[
    dict[str, dict[str, str]],
    dict[str, list[str]],
]:
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(
            f"Specification workbook not found: {EXCEL_PATH.resolve()}"
        )

    workbook = load_workbook(
        EXCEL_PATH,
        read_only=True,
        data_only=True,
    )

    if SPECIFICATION_SHEET not in workbook.sheetnames:
        available_sheets = ", ".join(workbook.sheetnames)

        raise ValueError(
            f"Sheet '{SPECIFICATION_SHEET}' was not found. "
            f"Available sheets: {available_sheets}"
        )

    worksheet = workbook[SPECIFICATION_SHEET]

    header_row = next(
        worksheet.iter_rows(
            min_row=1,
            max_row=1,
            values_only=True,
        )
    )

    headers = {
        normalize_text(value).lower(): index
        for index, value in enumerate(header_row)
        if value is not None
    }

    required_headers = {
        "scale",
        "layer label",
        "geoserver layer identifier",
        "service",
        "request type",
    }

    missing_headers = required_headers - set(headers)

    if missing_headers:
        workbook.close()

        raise ValueError(
            "The Specification sheet is missing these columns: "
            + ", ".join(sorted(missing_headers))
        )

    layers: dict[str, dict[str, str]] = {}
    expected_by_scale: dict[str, list[str]] = defaultdict(list)

    for row in worksheet.iter_rows(
        min_row=2,
        values_only=True,
    ):
        scale = normalize_text(
            row[headers["scale"]]
        )

        label = normalize_text(
            row[headers["layer label"]]
        )

        identifier = normalize_identifier(
            row[headers["geoserver layer identifier"]]
        )

        service = normalize_text(
            row[headers["service"]]
        )

        request_type = normalize_text(
            row[headers["request type"]]
        )

        # Ignore completely empty or incomplete specification rows.
        if not scale or not identifier:
            continue

        if not label:
            label = identifier

        if not service:
            service = "WMS"

        if not request_type:
            request_type = "GetMap"

        # Store each layer definition once.
        if identifier not in layers:
            layers[identifier] = {
                "label": label,
                "service": service,
                "request_type": request_type,
            }

        # Avoid adding the same identifier twice at one scale.
        if identifier not in expected_by_scale[scale]:
            expected_by_scale[scale].append(identifier)

    workbook.close()

    if not layers:
        raise ValueError(
            "No valid layer definitions were found in "
            f"the '{SPECIFICATION_SHEET}' sheet."
        )

    return layers, dict(expected_by_scale)


LAYERS, EXPECTED_LAYERS_BY_SCALE = read_specification()